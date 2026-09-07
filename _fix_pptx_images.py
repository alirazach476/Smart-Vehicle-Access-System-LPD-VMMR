"""Fix image sizes in fyp_ppt.pptx preserving aspect ratio and improving readability."""
from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

PPTX = Path(r"c:\Users\DELL\Downloads\FYP COMPLETE\FYP COMPLETE\fyp_ppt.pptx.images.bak")

# Content area below slide titles
MARGIN_L = 0.4
MARGIN_R = 0.4
MARGIN_T = 1.45
MARGIN_B = 0.25
GAP = 0.12

EMU = 914400


def inches(val) -> float:
    return val / EMU


def load_media_dims(pptx_path: Path) -> dict[str, tuple[int, int]]:
    dims: dict[str, tuple[int, int]] = {}
    with zipfile.ZipFile(pptx_path) as z:
        for name in z.namelist():
            if name.startswith("ppt/media/"):
                img = Image.open(io.BytesIO(z.read(name)))
                dims[name.split("/")[-1]] = img.size
    return dims


def load_slide_rels(pptx_path: Path, slide_num: int) -> str:
    rels_path = f"ppt/slides/_rels/slide{slide_num}.xml.rels"
    with zipfile.ZipFile(pptx_path) as z:
        return z.read(rels_path).decode("utf-8")


def media_for_shape(shape, rels: str, media_dims: dict[str, tuple[int, int]]) -> tuple[str, float]:
    blip = shape._element.find(".//{*}blip")
    rid = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed") if blip is not None else None
    target = None
    if rid:
        m = re.search(rf'Id="{rid}"[^>]*Target="([^"]+)"', rels)
        if not m:
            m = re.search(rf'Target="([^"]+)"[^>]*Id="{rid}"', rels)
        if m:
            target = m.group(1).replace("../media/", "")
    if target and target in media_dims:
        w, h = media_dims[target]
        return target, w / h
    return target or "unknown", shape.width / shape.height if shape.height else 1.0


def fit_in_cell(aspect: float, cell_w: float, cell_h: float) -> tuple[float, float]:
    if aspect <= 0:
        aspect = 1.0
    if cell_w / cell_h >= aspect:
        h = cell_h
        w = h * aspect
    else:
        w = cell_w
        h = w / aspect
    return w, h


def place_shape(shape, x: float, y: float, w: float, h: float) -> None:
    shape.left = Inches(x)
    shape.top = Inches(y)
    shape.width = Inches(w)
    shape.height = Inches(h)


def fit_banner(aspect: float, max_w: float, max_h: float) -> tuple[float, float]:
    w = max_w
    h = w / aspect
    if h > max_h:
        h = max_h
        w = h * aspect
    return w, h


def layout_grid(shapes, aspects, cols: int, rows: int, x0: float, y0: float, area_w: float, area_h: float) -> None:
    cell_w = (area_w - GAP * (cols - 1)) / cols
    cell_h = (area_h - GAP * (rows - 1)) / rows
    for idx, (shape, aspect) in enumerate(zip(shapes, aspects)):
        row, col = divmod(idx, cols)
        w, h = fit_in_cell(aspect, cell_w, cell_h)
        x = x0 + col * (cell_w + GAP) + (cell_w - w) / 2
        y = y0 + row * (cell_h + GAP) + (cell_h - h) / 2
        place_shape(shape, x, y, w, h)


def layout_slide(slide_num: int, slide, rels: str, media_dims: dict, slide_w: float, slide_h: float) -> int:
    pics = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    if not pics:
        return 0

    aspects = [media_for_shape(p, rels, media_dims)[1] for p in pics]
    content_w = slide_w - MARGIN_L - MARGIN_R
    content_h = slide_h - MARGIN_T - MARGIN_B
    x0, y0 = MARGIN_L, MARGIN_T

    if slide_num == 1:
        # Logo top-right, preserve aspect ratio
        aspect = aspects[0]
        max_w, max_h = 2.8, 2.8
        w, h = fit_in_cell(aspect, max_w, max_h)
        place_shape(pics[0], slide_w - MARGIN_R - w, 0.25, w, h)
        return 1

    if slide_num in {15, 18} and len(pics) == 2:
        bottom_h = content_h * 0.52
        banner_max_h = content_h - bottom_h - GAP
        banner_w, banner_h = fit_banner(aspects[0], content_w, banner_max_h)
        place_shape(pics[0], x0 + (content_w - banner_w) / 2, y0, banner_w, banner_h)

        bottom_y = y0 + banner_h + GAP
        w, h = fit_in_cell(aspects[1], content_w, bottom_h)
        place_shape(pics[1], x0 + (content_w - w) / 2, bottom_y + (bottom_h - h) / 2, w, h)
        return 2

    if slide_num == 11 and len(pics) == 3:
        bottom_h = content_h * 0.58
        banner_max_h = content_h - bottom_h - GAP
        banner_w, banner_h = fit_banner(aspects[0], content_w, banner_max_h)
        place_shape(pics[0], x0 + (content_w - banner_w) / 2, y0, banner_w, banner_h)

        bottom_y = y0 + banner_h + GAP
        cell_w = (content_w - GAP) / 2
        for i, pic in enumerate(pics[1:], start=1):
            w, h = fit_in_cell(aspects[i], cell_w, bottom_h)
            x = x0 + (i - 1) * (cell_w + GAP) + (cell_w - w) / 2
            y = bottom_y + (bottom_h - h) / 2
            place_shape(pic, x, y, w, h)
        return 3

    if slide_num in {21, 22} and len(pics) == 2:
        cell_w = (content_w - GAP) / 2
        for i, pic in enumerate(pics):
            w, h = fit_in_cell(aspects[i], cell_w, content_h)
            x = x0 + i * (cell_w + GAP) + (cell_w - w) / 2
            y = y0 + (content_h - h) / 2
            place_shape(pic, x, y, w, h)
        return 2

    if slide_num in {23, 24} and len(pics) == 1:
        w, h = fit_in_cell(aspects[0], content_w, content_h)
        place_shape(pics[0], x0 + (content_w - w) / 2, y0 + (content_h - h) / 2, w, h)
        return 1

    grid_map = {
        1: (1, 1),
        2: (2, 1),
        3: (3, 1),
        4: (2, 2),
        6: (3, 2),
        8: (4, 2),
    }
    cols, rows = grid_map.get(len(pics), (min(len(pics), 4), (len(pics) + 3) // 4))
    layout_grid(pics, aspects, cols, rows, x0, y0, content_w, content_h)
    return len(pics)


def main() -> None:
    media_dims = load_media_dims(PPTX)
    prs = Presentation(str(PPTX))
    slide_w = inches(prs.slide_width)
    slide_h = inches(prs.slide_height)

    total = 0
    for idx, slide in enumerate(prs.slides, 1):
        rels = load_slide_rels(PPTX, idx)
        count = layout_slide(idx, slide, rels, media_dims, slide_w, slide_h)
        if count:
            print(f"Slide {idx}: resized {count} image(s)")
            total += count

    target = Path(r"c:\Users\DELL\Downloads\FYP COMPLETE\FYP COMPLETE\fyp_ppt.pptx")
    temp_path = target.with_name("fyp_ppt_fixed.pptx")
    prs.save(str(temp_path))
    try:
        if target.exists():
            target.unlink()
        temp_path.replace(target)
        saved_as = target.name
    except OSError:
        saved_as = temp_path.name
        print(f"Note: close fyp_ppt.pptx in PowerPoint, then rename {saved_as} to fyp_ppt.pptx")
    print(f"Done. Updated {total} images in {saved_as}")


if __name__ == "__main__":
    main()
